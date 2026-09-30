"""Genera la presentación editable de 5 diapositivas (Jornada 1 – 1 de octubre).

Uso: python3 crear_presentacion.py  ->  ../02_presentacion_IA_CD_SINEP.pptx
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from contenido import (CARGOS, CREDITOS, CURSOS_ARTICULACION, JORNADA, POR_QUE_UNAD, PROGRAMA,
                       QUE_HACE, RUTA, SECTORES)
from estilos import OUT

NAVY, NAVY2 = RGBColor(0x0B, 0x2A, 0x5B), RGBColor(0x16, 0x3F, 0x7A)
GOLD, TEAL = RGBColor(0xFD, 0xB9, 0x13), RGBColor(0x19, 0xB3, 0xB1)
WHITE, INK, MUTED = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x1B, 0x24, 0x33), RGBColor(0x5B, 0x66, 0x78)
PAPER = RGBColor(0xF3, 0xF6, 0xFB)
FONT = "Montserrat"

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
IN = 914400


def i(x):
    return Emu(int(x * IN))


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(shape, i(x), i(y), i(w), i(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.1):
    """runs: str o lista de párrafos; cada párrafo str o lista de (texto, {size,color,bold})."""
    tb = s.shapes.add_textbox(i(x), i(y), i(w), i(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for k, para in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t
            f = r.font; f.name = FONT; f.size = Pt(o.get("size", size)); f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
    return tb


def pill(s, x, y, t, fill=GOLD, color=NAVY, size=12, w=None):
    w = w or (0.2 + len(t) * size * 0.0098)
    b = box(s, x, y, w, 0.36 * size / 12, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = b.text_frame; tf.margin_left = tf.margin_right = i(0.08); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = True
    r.font.color.rgb = color
    return w


def base(dark=False, num=None, tag=None, title=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.333, 7.5, NAVY if dark else PAPER)
    box(s, 0, 0, 0.14, 7.5, GOLD)
    if not dark:
        box(s, 0.14, 0, 13.2, 0.08, NAVY)
    fc = WHITE if dark else MUTED
    text(s, 0.6, 7.0, 6, 0.3, "UNAD · ECBTI · Del SINEP a la Ingeniería", size=10, color=fc, bold=True)
    if num:
        text(s, 11.7, 7.0, 1.1, 0.3, f"{num} / 5", size=10, color=fc, bold=True, align=PP_ALIGN.RIGHT)
    if tag:
        pill(s, 0.6, 0.5, tag, size=12)
    if title:
        text(s, 0.6, 0.98, 12, 0.8, title, size=32, color=WHITE if dark else NAVY, bold=True)
    return s


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# 1. Portada ---------------------------------------------------------------
s = base(dark=True)
for d, x, y, c in [(6.2, 8.4, -1.6, GOLD), (4.4, 9.5, -0.6, TEAL)]:
    o = box(s, x, y, d, d, None, c, MSO_SHAPE.OVAL); o.line.width = Pt(2)
pill(s, 0.8, 0.9, "Jornada 1 · Tecnologías Digitales", size=14)
text(s, 0.8, 1.7, 9.5, 3.2, [
    [("Ingeniería en", {"color": WHITE})],
    [("Inteligencia Artificial", {"color": GOLD})],
    [("y ", {"color": WHITE}), ("Ciencia de Datos", {"color": TEAL})]], size=50, bold=True, spacing=1.0)
text(s, 0.8, 4.75, 9, 0.5, "Escuela de Ciencias Básicas, Tecnología e Ingeniería – ECBTI · UNAD",
     size=16, color=WHITE)
x = 0.8
for a, b in [("100 %", "Virtual"), (str(CREDITOS), "Créditos"), ("Pregrado", "Título de ingeniero(a)")]:
    w = max(1.7, len(b) * 0.12 + 0.4)
    box(s, x, 5.45, w, 1.05, NAVY2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
    text(s, x + 0.2, 5.52, w - 0.3, 0.5, a, size=22, color=GOLD, bold=True)
    text(s, x + 0.2, 6.02, w - 0.3, 0.4, b, size=12, color=WHITE)
    x += w + 0.25
text(s, 7.8, 6.25, 5.0, 0.4, "Jueves 1 de octubre · 7:00 p. m.", size=14, color=WHITE, bold=True, align=PP_ALIGN.RIGHT)
notes(s, "Bienvenida (1 min). Presentar el programa: pregrado profesional 100 % virtual de 150 créditos, de la ECBTI. "
         "Pregunta rompehielo: ¿cómo creen que el celular reconoce su cara o una app sabe qué canción recomendar? "
         "Si es posible, reproducir aquí el video de 1:45.")

# 2. Perfil -----------------------------------------------------------------
s = base(num=2, tag="Perfil profesional", title="¿Qué hace este profesional?")
text(s, 0.6, 1.72, 12, 0.5, "Enseña a las máquinas a aprender y convierte los datos en decisiones que mejoran la vida de las personas.",
     size=15, color=MUTED)
for k, (e, t, d) in enumerate(QUE_HACE):
    x = 0.6 + k * 4.1
    box(s, x, 2.45, 3.85, 2.55, WHITE, RGBColor(0xDD, 0xE3, 0xEE), MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
    box(s, x, 2.45, 3.85, 0.1, [GOLD, TEAL, NAVY2][k])
    text(s, x + 0.3, 2.75, 1, 0.7, e, size=30)
    text(s, x + 0.3, 3.45, 3.3, 0.75, t, size=16, color=NAVY, bold=True)
    text(s, x + 0.3, 4.2, 3.3, 0.8, d[0].upper() + d[1:], size=13, color=INK)
box(s, 0.6, 5.3, 12.1, 1.45, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
text(s, 0.95, 5.45, 3.2, 1.2, [[("¿Es para ti?", {"color": GOLD, "size": 20, "bold": True})],
                               [("Perfil de ingreso", {"color": WHITE, "size": 12})]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 4.0, 5.45, 8.5, 1.2,
     "Te gusta la tecnología, resolver problemas y pensar con lógica. Tienes curiosidad, creatividad, "
     "disposición para trabajar en equipo y aprender de forma autónoma. No necesitas experiencia previa.",
     size=13, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Explicar con ejemplos reales: detectar fraudes en un banco, predecir la demanda de productos en una tienda, "
         "apoyar diagnósticos médicos con imágenes, chatbots que responden en lenguaje natural. "
         "Recalcar la ética y el compromiso social (perfil de egreso del Documento Maestro).")

# 3. Campo laboral ------------------------------------------------------------
s = base(num=3, tag="Campo laboral", title="¿En qué puedes trabajar?")
for k, (e, t) in enumerate(CARGOS):
    x, y = 0.6 + (k % 2) * 6.1, 1.95 + (k // 2) * 0.95
    box(s, x, y, 5.9, 0.8, WHITE, RGBColor(0xDD, 0xE3, 0xEE), MSO_SHAPE.ROUNDED_RECTANGLE, 0.2)
    text(s, x + 0.2, y, 0.6, 0.8, e, size=20, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.85, y, 4.9, 0.8, t, size=14, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 4.95, 6, 0.4, "Sectores que necesitan tu talento", size=16, color=NAVY, bold=True)
x, y = 0.6, 5.45
for t in SECTORES:
    w = 0.3 + len(t) * 0.118
    if x + w > 12.7:
        x, y = 0.6, y + 0.52
    pill(s, x, y, t, fill=NAVY, color=WHITE, size=12, w=w)
    x += w + 0.15
notes(s, "Contar qué problemas resuelve en una institución: reducir tiempos de atención con automatización, "
         "anticipar deserción estudiantil, optimizar rutas de distribución, tableros de datos para alcaldías. "
         "Mencionar que también puede emprender o trabajar de forma remota para empresas nacionales e internacionales.")

# 4. Articulación ----------------------------------------------------------------
s = base(num=4, tag="Articulación SINEP – ECBTI", title="Empieza la universidad desde el colegio")
total = sum(c for _, c, _ in CURSOS_ARTICULACION)
for k, (n, c, d) in enumerate(CURSOS_ARTICULACION):
    x = 0.6 + k * 2.75
    box(s, x, 1.95, 2.55, 2.9, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, x + 0.25, 2.1, 2.1, 0.8, [[(str(c), {"size": 40, "color": GOLD, "bold": True}),
                                         ("  créditos", {"size": 12, "color": GOLD, "bold": True})]])
    text(s, x + 0.25, 2.95, 2.1, 0.8, n, size=15, color=WHITE, bold=True)
    text(s, x + 0.25, 3.8, 2.1, 1.0, d, size=11, color=WHITE)
text(s, 0.6, 5.05, 8.2, 0.5, [[("✔ Hasta ", {}), (f"{total} créditos", {"color": NAVY, "bold": True}),
                               (" reconocidos al ingresar al programa", {})]], size=15, color=INK)
text(s, 0.6, 5.5, 8.2, 1.2, ["✔ Ahorras tiempo y dinero   ✔ Continuidad académica",
                             "✔ Mayor posibilidad de ingreso   ✔ Fortaleces tu proyecto de vida"], size=13, color=MUTED)
box(s, 9.0, 1.95, 3.75, 4.8, WHITE, RGBColor(0xDD, 0xE3, 0xEE), MSO_SHAPE.ROUNDED_RECTANGLE, 0.06)
text(s, 9.25, 2.05, 3.3, 0.4, "Tu ruta, paso a paso", size=14, color=NAVY, bold=True)
for k, (n, t, _) in enumerate(RUTA):
    y = 2.55 + k * 0.68
    o = box(s, 9.25, y, 0.46, 0.46, GOLD, shape=MSO_SHAPE.OVAL)
    tf = o.text_frame; tf.margin_left = tf.margin_right = 0
    r = tf.paragraphs[0].add_run(); r.text = n; r.font.size = Pt(13); r.font.bold = True
    r.font.color.rgb = NAVY; r.font.name = FONT; tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(s, 9.85, y - 0.02, 2.8, 0.5, t, size=13, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Explicar: qué es la articulación y el reconocimiento de saberes; qué significa aprobar un curso desde la "
         "formación media; cuántos créditos se reconocen; cómo matricular con la coordinación del SINEP. "
         "IMPORTANTE: confirmar con la coordinación SINEP–ECBTI la lista oficial de cursos y créditos antes de presentar.")

# 5. ¿Por qué la UNAD? -------------------------------------------------------------
s = base(dark=True, num=5, tag="¿Por qué estudiar en la UNAD?", title="Tu futuro empieza hoy")
for k, (e, t, d) in enumerate(POR_QUE_UNAD):
    x, y = 0.6 + (k % 3) * 4.1, 1.95 + (k // 3) * 1.75
    if k >= 3:
        x += 2.05
    box(s, x, y, 3.85, 1.55, NAVY2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, x + 0.25, y + 0.18, 0.7, 0.6, e, size=24)
    text(s, x + 0.95, y + 0.15, 2.75, 0.6, t, size=13, color=GOLD, bold=True)
    text(s, x + 0.95, y + 0.78, 2.75, 0.75, d, size=10.5, color=WHITE)
box(s, 0.6, 5.7, 12.15, 1.05, GOLD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
text(s, 0.9, 5.7, 8.0, 1.05, [[("📝 Llena el formulario de interés ", {"bold": True, "size": 18}),
                               ("y elige IA y Ciencia de Datos", {"size": 14})]],
     color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
text(s, 8.5, 5.7, 4.0, 1.05, "Del SINEP a la Ingeniería", size=18, color=NAVY, bold=True, align=PP_ALIGN.RIGHT,
     anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Cerrar con los beneficios (ahorro de tiempo y económico, continuidad, movilidad, proyecto de vida). "
         "Invitar a diligenciar el formulario y abrir la sesión de preguntas. "
         "Recordar el espacio de orientación del 29 de octubre.")

out = OUT / "02_presentacion_IA_CD_SINEP.pptx"
prs.save(out)
print(out)
